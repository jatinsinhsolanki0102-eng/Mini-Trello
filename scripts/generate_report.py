"""
Generate the Mini-Trello project report as a DOCX file.

    python scripts/generate_report.py

Output:  Mini-Trello-Report.docx  (in the project root)

Embedded screenshots come from the screenshots/ folder (real captures).
"""

import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, "screenshots")
OUT = os.path.join(ROOT, "Mini-Trello-Report.docx")

PURPLE = RGBColor(0x7C, 0x3A, 0xED)
GRAY = RGBColor(0x59, 0x59, 0x59)


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.color.rgb = PURPLE
    return p


def h2(doc, text):
    return doc.add_heading(text, level=2)


def para(doc, text, style=None):
    return doc.add_paragraph(text, style=style)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def cover_table(doc):
    table = doc.add_table(rows=6, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [
        ("Project", "Mini-Trello Kanban Board"),
        ("Student Name", "Jatinsinh Solanki"),
        ("Enrollment Number", "230390116025"),
        ("Team", "Individual"),
        ("Semester", "7th Semester"),
        ("Methodology", "Agile / Scrum (two one-week sprints)"),
    ]
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].paragraphs[0].runs[0].bold = True


def api_table(doc):
    table = doc.add_table(rows=5, cols=4)
    table.style = "Light Grid Accent 1"
    headers = ["Method", "Endpoint", "Purpose", "Status"]
    for j, htext in enumerate(headers):
        table.rows[0].cells[j].text = htext
    data = [
        ("GET", "/api/tasks", "Fetch all tasks", "200 OK"),
        ("POST", "/api/tasks", "Create task (status = todo)", "201 Created"),
        ("PUT", "/api/tasks/:id", "Update status / fields", "200 OK"),
        ("DELETE", "/api/tasks/:id", "Delete a task", "200 OK"),
    ]
    for i, row in enumerate(data, start=1):
        for j, value in enumerate(row):
            table.rows[i].cells[j].text = value


def test_table(doc):
    table = doc.add_table(rows=12, cols=6)
    table.style = "Light Grid Accent 1"
    headers = ["#", "Test", "Input/Action", "Expected", "Actual", "Status"]
    for j, htext in enumerate(headers):
        table.rows[0].cells[j].text = htext
    data = [
        ("1", "GET empty DB", "GET /api/tasks", "200, empty list", "200, []", "Pass"),
        ("2", "GET multiple", "Create 2, GET", "200, both", "200, both", "Pass"),
        ("3", "POST task", "Valid title+desc", "201, todo", "201, todo", "Pass"),
        ("4", "POST empty title", 'title = ""', "400 Title required", "400", "Pass"),
        ("5", "POST empty desc", 'description = ""', "400 Desc required", "400", "Pass"),
        ("6", "PUT status", '{"status":"in_progress"}', "200, updated", "200", "Pass"),
        ("7", "PUT invalid status", '{"status":"review"}', "400", "400", "Pass"),
        ("8", "PUT missing task", "id = 99999", "404", "404", "Pass"),
        ("9", "DELETE task", "existing id", "200, removed", "200, removed", "Pass"),
        ("10", "DELETE missing", "id = 99999", "404", "404", "Pass"),
        ("11", "Transitions", "todo→ip→done→todo", "persist", "all pass", "Pass"),
    ]
    for i, row in enumerate(data, start=1):
        for j, value in enumerate(row):
            table.rows[i].cells[j].text = value


def screenshot(doc, name, caption):
    path = os.path.join(SHOTS, name)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(6.0))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.left_margin = section.right_margin = Inches(1.0)

    # ---- Cover ----
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Mini-Trello Kanban Board")
    run.font.size = Pt(28)
    run.font.color.rgb = PURPLE
    run.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Task Management Application — Project Report").italic = True

    doc.add_paragraph()
    cover_table(doc)
    doc.add_page_break()

    # ---- Sections ----
    h1(doc, "1. Project Overview")
    para(doc, "Mini-Trello is a simple, modern, responsive single-page Kanban task-"
              "management application. A user can create tasks, view them grouped into "
              "three status columns (To Do, In Progress, Done), move tasks between the "
              "columns with Next/Previous controls, and delete tasks. All data is stored "
              "in a real database and every action goes through a RESTful API, so the "
              "board state survives browser refreshes.")
    para(doc, "The project was developed to practice the full stack: a React frontend, "
              "a Flask backend, and a relational database (MySQL; SQLite is used locally "
              "for zero-setup development), while applying Agile/Scrum across two "
              "one-week sprints.")

    h1(doc, "2. Objectives")
    bullets(doc, [
        "Build a task-management board with three columns and create/move/delete flows.",
        "Implement a RESTful API (GET, POST, PUT, DELETE) with JSON in/out.",
        "Store tasks in a real database with persistence across restarts.",
        "Integrate a React frontend with a Flask backend over HTTP.",
        "Apply Agile/Scrum methodology with two one-week sprints.",
    ])

    h1(doc, "3. Technology Stack")
    table = doc.add_table(rows=6, cols=2)
    table.style = "Light Grid Accent 1"
    data = [
        ("Frontend", "React 19, JavaScript, CSS custom properties (design tokens), Vite 6"),
        ("Backend", "Python 3.13, Flask 3.1, Flask-CORS, SQLAlchemy 2"),
        ("Database", "MySQL 8 (schema provided); SQLite as zero-setup local default"),
        ("UI / Theming", "Dark purple signature theme + Light theme, persisted in localStorage"),
        ("API", "REST, JSON"),
        ("Testing", "Pytest (26 cases), Playwright (browser E2E + screenshots)"),
    ]
    for i, (k, v) in enumerate(data):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        table.rows[i].cells[1].text = v

    h1(doc, "4. System Architecture")
    para(doc, "Browser (React SPA) → REST API → Flask backend → Database", style="No Spacing")
    code = doc.add_paragraph("Browser (React SPA, Vite)  :5173\n"
                             "        │ fetch()\n"
                             "        ▼\n"
                             "REST API  :5173/api/*   (Vite dev proxy → Flask :5000)\n"
                             "        │\n"
                             "        ▼\n"
                             "Flask backend  (CORS, routes, validation, error handlers)\n"
                             "        │ SQLAlchemy ORM (parameterized queries)\n"
                             "        ▼\n"
                             "Database  (tasks table — SQLite / MySQL)")
    for run in code.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)

    h1(doc, "5. Database Design")
    para(doc, "The application uses a single tasks table:")
    table = doc.add_table(rows=7, cols=3)
    table.style = "Light Grid Accent 1"
    cols = ["Column", "Type", "Notes"]
    for j, c in enumerate(cols):
        table.rows[0].cells[j].text = c
    rows = [
        ("id", "INT (PK, auto-increment)", "unique task id"),
        ("title", "VARCHAR(200)", "required"),
        ("description", "TEXT", "required"),
        ("status", "VARCHAR(20)", "todo / in_progress / done (CHECK constraint)"),
        ("created_at", "DATETIME", "set automatically"),
        ("updated_at", "DATETIME", "updated on change"),
    ]
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            table.rows[i].cells[j].text = value
    para(doc, "The MySQL creation script is database/schema.sql. Tables are created "
              "automatically by the backend on startup.")

    h1(doc, "6. UI Design & Theming")
    para(doc, "The final UI is a modern purple-based design system. Every colour and "
              "shadow is defined once as a CSS custom property (design token) in "
              "frontend/src/styles/styles.css, so the whole app is themed consistently "
              "and can be re-skinned without touching component code.")
    bullets(doc, [
        "Dark Mode (default / signature): purple-black background (#0F0B1A), raised "
        "purple surfaces, high-contrast text, and purple (#7C3AED) accent buttons.",
        "Light Mode: soft lavender background (#F7F5FC) with white cards and the same "
        "purple accent colour, fully readable, same layout with zero layout shift.",
        "Segmented theme selector (Light | Dark) in the header; choice saved under "
        "mini-trello-theme in localStorage and restored before first paint (no flash "
        "of the wrong theme).",
        "Colour-coded status accents for the three columns (purple for To Do, violet "
        "for In Progress, green for Done), matching status badges on every card.",
        "Rounded icon buttons, soft shadows, subtle entrance animations that respect "
        "prefers-reduced-motion, and responsive breakpoints at 960/640/420 px.",
    ])
    para(doc, "Theme switching is fully covered by automated browser checks: default "
              "dark, switch to light, persistence after a refresh, then switch back "
              "to dark.")

    h1(doc, "7. User Stories & Product Backlog")
    bullets(doc, [
        "US1 (High, 5 pts): create a new task and add it to the board.",
        "US2 (High, 3 pts): see all tasks grouped by status.",
        "US3 (High, 5 pts): move tasks To Do → In Progress → Done.",
        "US4 (Medium, 3 pts): delete a task with confirmation.",
    ])
    para(doc, "Total estimate: 16 story points. Full acceptance criteria in "
              "docs/PRODUCT_BACKLOG.md. Sprint 1 owns US1 + US2 (foundation); "
              "Sprint 2 owns US3 + US4 plus full integration.")

    h1(doc, "8. Sprint 1 — Foundation & Setup")
    bullets(doc, [
        "Goal: working database, GET/POST APIs, static board UI.",
        "Designed the tasks schema and created the database.",
        "Implemented the Task model with status validation and a DB constraint.",
        "Implemented and tested GET /api/tasks and POST /api/tasks.",
        "Scaffolded the React + Vite frontend with the three-column board UI.",
        "Deliverable: working database, GET/POST APIs testable via Postman, static UI.",
    ])

    h1(doc, "9. Sprint 2 — Integration & Delivery")
    bullets(doc, [
        "Goal: complete integration and deliver the product.",
        "Implemented PUT /api/tasks/:id and DELETE /api/tasks/:id with 404 handling.",
        "Connected the frontend through a centralized API service.",
        "Implemented create, move (Next/Previous), and delete with confirmation.",
        "Added loading states, error banners, frontend validation and responsive polish.",
        "Redesigned the final UI to the purple design system with Light/Dark themes and localStorage persistence.",
        "Wrote and ran 26 pytest cases and 23 browser E2E checks; captured screenshots in both themes.",
        "Prepared README, this report, and the source ZIP.",
    ])

    h1(doc, "10. REST API Documentation")
    api_table(doc)
    para(doc, "BASE URL: http://localhost:5000/api. Errors use {\"error\": \"...\"} with "
              "status 400 (validation), 404 (missing task), 500 (server).")
    para(doc, "Sample POST request:", style="No Spacing")
    para(doc, '{"title": "Create UI", "description": "Build the Mini-Trello interface"}',
         style="No Spacing")
    para(doc, "Sample POST response (201):", style="No Spacing")
    para(doc, '{"id": 1, "title": "Create UI", "description": "Build the Mini-Trello interface", "status": "todo"}',
         style="No Spacing")

    h1(doc, "11. Testing")
    para(doc, "Backend tests (pytest):")
    test_table(doc)
    para(doc, 'Command: python -m pytest -q  →  26 passed.')
    para(doc, "End-to-end verification: live API create/move/delete/refresh, "
              "full server-restart persistence test, and real-browser flows via "
              "Playwright (create modal, movement, inline delete confirmation, "
              "mobile layout, and theme switching with persistence). "
              "Result: 23/23 E2E checks passed.")

    h1(doc, "12. Screenshots")
    para(doc, "All screenshots below were captured from the running application "
              "using scripts/capture_screenshots.py (real browser, live data, "
              "light and dark themes).")
    screenshot(doc, "01-board-light.png", "Figure 1 — Main board (Light theme)")
    screenshot(doc, "02-board-dark.png", "Figure 2 — Main board (Dark theme, signature)")
    screenshot(doc, "03-modal-light.png", "Figure 3 — Create New Task modal (Light)")
    screenshot(doc, "04-modal-dark.png", "Figure 4 — Create New Task modal filled (Dark)")
    screenshot(doc, "05-task-in-todo.png", "Figure 5 — Task in To Do")
    screenshot(doc, "06-task-in-progress.png", "Figure 6 — Task in In Progress")
    screenshot(doc, "07-task-in-done.png", "Figure 7 — Task in Done")
    screenshot(doc, "08-delete-confirmation.png", "Figure 8 — Delete confirmation")
    screenshot(doc, "09-database-table.png", "Figure 9 — tasks table (live data)")
    screenshot(doc, "10-mobile-dark.png", "Figure 10 — Responsive mobile layout (Dark)")
    screenshot(doc, "11-api-testing.png", "Figure 11 — REST API requests/responses (live)")

    h1(doc, "13. Agile Execution Summary")
    para(doc, "Two one-week sprints with planning, daily stand-ups, a review and a "
              "retrospective. All ceremony details, the daily log and the responsibility "
              "table are in the docs/ folder (individual project — Jatinsinh Solanki, "
              "Enrollment 230390116025).")

    h1(doc, "14. Future Improvements")
    bullets(doc, [
        "Drag-and-drop between columns; in-place card editing.",
        "Optional assigned_to field and assignee filters.",
        "Search, filters and pagination.",
        "Deployment to free hosting platforms.",
    ])

    doc.save(OUT)
    print("Report saved to", OUT)


if __name__ == "__main__":
    build()